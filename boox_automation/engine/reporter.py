"""结果回写器：将执行结果写回 Excel 的「测试结果」列。"""
from __future__ import annotations

import logging
from pathlib import Path

import openpyxl

from boox_automation.engine.parser import ParsedCase

logger = logging.getLogger(__name__)

from boox_automation.core.config import case_column

COL_RESULT = case_column("result")
COL_REMARK = case_column("remark")


class ExcelReporter:
    """将执行结果写回 Excel 文件。"""

    def __init__(self, filepath: str):
        self.filepath = Path(filepath)
        if not self.filepath.exists():
            raise FileNotFoundError(f"Excel 文件不存在: {filepath}")
        self._wb = openpyxl.load_workbook(filepath)

    def write_results(self, sheet_name: str, cases: list[ParsedCase]) -> None:
        """将用例执行结果写回指定 sheet。"""
        ws = self._wb[sheet_name]

        summary = {"pass": 0, "fail": 0, "skip": 0}

        for case in cases:
            row = case.row_number
            status = self._case_status(case)
            summary[status] = summary.get(status, 0) + 1

            cell = ws.cell(row=row, column=COL_RESULT + 1)  # openpyxl 1-indexed
            status_map = {
                "pass": "通过",
                "fail": "不通过",
                "skip": "跳过",
            }
            cell.value = status_map.get(status, "")

            # 测试备注：写详细步骤 + 跳过原因
            remark = self._build_detail(case)
            if remark:
                ws.cell(row=row, column=COL_REMARK + 1).value = remark

        logger.info(f"{sheet_name}: 通过={summary['pass']}, 不通过={summary['fail']}, 跳过={summary['skip']}")

    @staticmethod
    def _case_status(case: ParsedCase) -> str:
        """根据步骤状态和预期结果状态判定用例整体结果。"""
        # 前置条件不满足 → 跳过
        if case.skip_reason:
            return "skip"

        statuses = {s.status for s in case.steps}
        # 预期结果失败也计入
        er_statuses = {ep.status for ep in case.expected_pages if ep.status}
        all_statuses = statuses | er_statuses

        # "" 表示步骤未执行（setup 失败或异常中断）→ 失败
        if "fail" in all_statuses or "" in statuses:
            return "fail"
        if not all_statuses or all_statuses == {"skip"}:
            return "skip"
        return "pass"

    @staticmethod
    def _build_detail(case: ParsedCase) -> str:
        """生成步骤执行详情（含前置条件跳过原因）。"""
        lines = []

        if case.skip_reason:
            lines.append(f"✗ {case.skip_reason}")

        for s in case.steps:
            if s.status == "skip" and not s.tag:
                continue
            if s.status == "skip" and case.skip_reason:
                continue  # 前置条件导致的批量 skip 不重复展示
            icon = {"pass": "✓", "fail": "✗", "skip": "-"}.get(s.status, "?")
            lines.append(f"{icon} 步骤{s.seq}: {s.action}【{s.tag}】→ {s.element_key or '未匹配'}")

        # 预期结果校验详情
        for ep in case.expected_pages:
            if not ep.status:
                continue
            icon = {"pass": "✓", "fail": "✗", "skip": "-"}.get(ep.status, "?")
            label = f"步骤{ep.step_seq}" if ep.step_seq else "最终"
            detail = f"预期结果: 【{ep.tag}】→ {ep.expected_key or '未匹配'}"
            lines.append(f"{icon} [{label}] {detail}")

        return "\n".join(lines)

    def save(self, output_path: str | None = None) -> None:
        """保存 Excel 文件。"""
        target = output_path or self.filepath
        self._wb.save(str(target))
        logger.info(f"结果已保存: {target}")
