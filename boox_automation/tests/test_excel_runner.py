"""Excel 驱动测试 — 集成到 pytest 框架。

每条 Excel 用例作为独立的 pytest 用例执行，自动获得 note_test_initial 的完整环境清理。

运行方式:
    pytest boox_automation/tests/test_excel_runner.py -s
"""
from __future__ import annotations
import sys
import os
import logging
from pathlib import Path

# 确保项目根目录在 sys.path 中（兼容直接 python 运行及 VS Code 等不以项目根为 cwd 的方式）
_project_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _project_root not in sys.path:
    sys.path.insert(0, _project_root)

import pytest
import openpyxl
from boox_automation.driver import driver
from boox_automation.ui_ops.operations import (
    Operation_method, set_step_context, clear_step_context, clear_element_ctx,
)
from boox_automation.tests.helpers import Public_method
from boox_automation.engine.parser import (
    parse_steps, parse_preconditions, parse_expected_results,
    ParsedCase, ExpectedPageRef, check_conditions, validate_case,
)
from boox_automation.engine.elements import ElementMatcher
from boox_automation.engine.result_store import set_cases

logger = logging.getLogger(__name__)

# ============================================================
from boox_automation.core.config import excel_priority_filter, case_path, excel_source, test_modules as cfg_test_modules, case_column

EXCEL_FILE = case_path()
PRIORITY_FILTER = excel_priority_filter()
_PRIORITY_ORDER = {"P0": 0, "P1": 1, "P2": 2, "TEST": 3}
# ============================================================


def _get_device_info_safe() -> dict:
    """安全获取设备信息，失败返回空字典（不影响模块加载）。"""
    try:
        from boox_automation.devices.info import Device_basic_information
        dbi = Device_basic_information()
        return dbi.get_device_info() or {}
    except Exception:
        return {}


def _get_device_id_safe() -> str:
    """安全获取当前连接的设备ID，失败返回空字符串。"""
    try:
        from boox_automation.devices.info import Device_basic_information
        dbi = Device_basic_information()
        return dbi.get_connected_device_ids() or ""
    except Exception:
        return ""


_element_matcher = ElementMatcher()


def _parse_case_rows(rows: list[list[str]]) -> list[ParsedCase]:
    """从二维数组解析用例文本（不含元素匹配，只做文本解析）。"""
    cases = []

    for row_idx_0, row in enumerate(rows):
        if row_idx_0 == 0:
            continue  # 跳过表头行
        row1 = row_idx_0 + 1  # 1-based row number

        def _col(c): return row[c] if c < len(row) else ""

        title = _col(case_column("title"))
        steps_text = _col(case_column("steps"))
        priority = _col(case_column("priority"))
        module = _col(case_column("module"))
        precondition_text = _col(case_column("precondition"))
        expected_text = _col(case_column("expected"))

        if not title or not steps_text:
            continue

        case = ParsedCase(
            title=str(title).strip(),
            priority=str(priority).strip() if priority else "",
            module=str(module).strip() if module else "",
            row_number=row1,
            steps=parse_steps(str(steps_text).strip()),
            preconditions=parse_preconditions(precondition_text),
            expected_pages=parse_expected_results(expected_text),
        )
        cases.append(case)

    return cases


def _resolve_case_elements(cases: list[ParsedCase]) -> list[ParsedCase]:
    """为过滤后的用例匹配元素（仅对要执行的少量用例做匹配）。"""
    from boox_automation.engine.elements import get_element_loader
    _loader = get_element_loader()

    for case in cases:
        ctx = f"R{case.row_number} [{case.priority}] {case.title}"
        for step in case.steps:
            if step.tag:
                ek = _element_matcher.match(
                    step.tag, page_context=case.module, case_context=ctx, warn=None)
                if ek:
                    step.element_key = ek
                    if not step.action or step.action == "click":
                        yaml_action = _element_matcher.suggest_action(ek)
                        if yaml_action:
                            step.action = yaml_action

        for ep in case.expected_pages:
            if ep.tag:
                ep.expected_key = _loader.match_expected_result(ep.tag)

        # 收集阶段诊断：预期结果匹配失败的 tag 提前暴露
        unresolved = [
            ep for ep in case.expected_pages
            if ep.tag and not ep.expected_key
        ]
        if unresolved:
            diag = _loader.get_expected_diagnostics()
            total = diag.get("total_results", 0)
            if total == 0:
                logger.warning(
                    f"R{case.row_number}「{case.title}」"
                    f"I列预期结果 {', '.join(f'【{ep.tag}】' for ep in unresolved)} "
                    f"未匹配：预期结果工作表不存在，用例执行时将报错"
                )
            else:
                index_keys = sorted(diag.get("match_index", {}).keys())
                logger.warning(
                    f"R{case.row_number}「{case.title}」"
                    f"I列预期结果 {', '.join(f'【{ep.tag}】' for ep in unresolved)} "
                    f"在预期结果工作表 B 列中未匹配。"
                    f"当前 B 列已有: {', '.join(index_keys) if index_keys else '(空)'}"
                )

    # 收集阶段输出预期结果加载摘要（仅跳过详情，汇总见模块级日志）
    diag = _loader.get_expected_diagnostics()
    skipped_list = diag.get("skipped", [])
    if skipped_list:
        for s in skipped_list:
            logger.warning(
                f"  ↳ 跳过 第{s['row']}行【{s['key']}】: {s['reason']}"
            )

    return cases


def _apply_filters(cases: list[ParsedCase], priority_spec: str) -> list[ParsedCase]:
    """优先级筛选 + 模块筛选。"""
    spec = priority_spec.upper().strip()
    if spec == "P0":
        cases = [c for c in cases if c.priority == "P0"]
    elif spec == "P1":
        cases = [c for c in cases if c.priority in ("P1", "P2")]
    elif spec == "P2":
        cases = [c for c in cases if c.priority == "P2"]
    elif spec == "TEST":
        cases = [c for c in cases if c.priority.upper() == "TEST"]

    modules_filter = cfg_test_modules()
    if modules_filter:
        cases = [c for c in cases if c.module in modules_filter]

    return cases


def _load_cases(excel_path: str, sheets: list[str], priority_spec: str) -> list[ParsedCase]:
    """按配置的 excel_source 加载用例（单一路径，不回退）。

    NOTE_TEST_CASE_PATH 可覆盖用例文件路径（与 excel_source 独立），
    便于试点批次：元素/预期走 cache，用例走本地 pilot xlsx。
    """
    env_case = os.environ.get("NOTE_TEST_CASE_PATH", "").strip()
    if env_case and Path(env_case).is_file():
        logger.info(f"用例加载: NOTE_TEST_CASE_PATH 覆盖 — {env_case}")
        sheet_data = _read_local_sheets(env_case, sheets)
        return _parse_cases_from_sheets(sheet_data, sheets, priority_spec)

    source = excel_source()

    if source == "local":
        logger.info(f"用例加载: 本地 Excel — {excel_path}")
        sheet_data = _read_local_sheets(excel_path, sheets)
        return _parse_cases_from_sheets(sheet_data, sheets, priority_spec)

    if source == "cache":
        logger.info("用例加载: 本地缓存")
        cached = _load_cases_from_cache(sheets, priority_spec)
        if cached is not None:
            return cached
        raise RuntimeError(
            "用例加载失败: excel.source=cache，但无可用缓存。"
            "请先以 cloud 模式运行一次生成缓存，或改为 local 模式。"
        )

    if source == "cloud":
        from boox_automation.core.feishu import check_feishu_reachable, save_cache
        logger.info(f"用例加载: 飞书云端（目标工作表: {', '.join(sheets)}）")
        if not check_feishu_reachable():
            raise RuntimeError(
                "用例加载失败: excel.source=cloud，但飞书 API 不可达。"
                "请检查网络连接或切换为 cache/local 模式。"
            )
        try:
            sheet_data = _read_cloud_sheets(sheets)
            save_cache({"sheets": sheet_data}, "test_cases")
            logger.info("用例来源: 飞书云端（已更新本地缓存）")
            return _parse_cases_from_sheets(sheet_data, sheets, priority_spec)
        except Exception as e:
            raise RuntimeError(
                f"用例加载失败: excel.source=cloud，飞书云端加载异常。"
                f"错误: {e}"
            ) from e

    raise RuntimeError(f"未知的 excel.source: {source}")


def _read_local_sheets(excel_path: str, sheets: list[str]) -> dict[str, list[list[str]]]:
    """从本地 Excel 读取指定 sheet 的原始行数据。"""
    wb = openpyxl.load_workbook(excel_path)
    sheet_data = {}
    for sheet in sheets:
        if sheet not in wb.sheetnames:
            logger.warning(f"本地工作表「{sheet}」不存在，跳过")
            continue
        ws = wb[sheet]
        rows = []
        for row in ws.iter_rows(min_row=1, values_only=True):
            rows.append([str(v) if v is not None else "" for v in row])
        sheet_data[sheet] = rows
    wb.close()
    return sheet_data


def _read_cloud_sheets(sheets: list[str]) -> dict[str, list[list[str]]]:
    """从飞书云端读取指定 sheet 的原始行数据。"""
    from boox_automation.core.feishu import read_sheet_by_name
    from boox_automation.core.config import feishu_test_case_token
    token = feishu_test_case_token()
    sheet_data = {}
    for sheet in sheets:
        rows = read_sheet_by_name(sheet, token)
        if rows:
            sheet_data[sheet] = rows
    return sheet_data


def _parse_cases_from_sheets(sheet_data: dict, sheets: list[str],
                              priority_spec: str) -> list[ParsedCase]:
    """从原始行数据解析用例列表（解析→筛选→匹配元素）。"""
    all_cases = []
    for sheet in sheets:
        rows = sheet_data.get(sheet, [])
        if not rows:
            logger.warning(f"工作表「{sheet}」无数据")
            continue
        # 1. 文本解析（快速，无元素匹配）
        sheet_cases = _parse_case_rows(rows)
        logger.debug(f"工作表「{sheet}」: 阶段1 文本解析 → {len(sheet_cases)} 条用例")
        # 2. 优先级 + 模块筛选
        before_filter = len(sheet_cases)
        sheet_cases = _apply_filters(sheet_cases, priority_spec)
        filtered = before_filter - len(sheet_cases)
        if filtered:
            logger.debug(
                f"工作表「{sheet}」: 阶段2 筛选 → 过滤 {filtered} 条，"
                f"剩余 {len(sheet_cases)} 条（优先级={priority_spec}）"
            )
        else:
            logger.debug(f"工作表「{sheet}」: 阶段2 筛选 → 无需过滤，保留全部 {len(sheet_cases)} 条")
        # 3. 只对筛选后的用例做元素匹配
        sheet_cases = _resolve_case_elements(sheet_cases)
        executable = sum(1 for c in sheet_cases if any(s.element_key for s in c.steps))
        logger.info(f"工作表「{sheet}」: {len(sheet_cases)} 条用例（{executable} 条可执行）")
        all_cases.extend(sheet_cases)
    return all_cases


def _load_cases_from_cache(sheets: list[str], priority_spec: str) -> list | None:
    """从缓存加载用例，缓存过期或无数据返回 None。"""
    from boox_automation.core.feishu import load_cache
    cached = load_cache("test_cases")
    if not cached:
        return None
    sheet_data = cached.get("sheets", {})
    if not sheet_data:
        return None
    cases = _parse_cases_from_sheets(sheet_data, sheets, priority_spec)
    return cases if cases else None


def _make_case_id(case: ParsedCase) -> str:
    prefix = "[跳过]" if case.skip_reason else ""
    return f"{prefix}R{case.row_number}-{case.title[:30]}"


# ── 用例收集 ──
from boox_automation.core.config import test_case_sheets as _test_case_sheets
_sheets = _test_case_sheets()
logger.debug(
    f"用例收集配置: 目标工作表={_sheets}, "
    f"优先级筛选={PRIORITY_FILTER}, "
    f"数据源={excel_source()}"
)
logger.info("── 开始收集用例 ──")
_all_cases = _load_cases(EXCEL_FILE, _sheets, PRIORITY_FILTER)
set_cases(_all_cases)
_executable = [c for c in _all_cases if any(s.element_key for s in c.steps)]
_executable.sort(key=lambda c: (_PRIORITY_ORDER.get(c.priority, 99), c.row_number))
_skipped = len(_all_cases) - len(_executable)

# 收集完成：输出汇总（用例 + 元素 + 预期结果）
from boox_automation.engine.elements import get_element_loader as _get_loader_for_summary
_summary_loader = _get_loader_for_summary()
_elem_count = len(_summary_loader)
_er_diag = _summary_loader.get_expected_diagnostics()
_er_total = _er_diag.get("total_results", 0)
_er_skipped = len(_er_diag.get("skipped", []))
logger.info(
    "── 收集完成: %d 条用例（%d 可执行, %d 跳过）| "
    "元素 %d 个 | 预期结果 %d 条（跳过 %d 条）──",
    len(_all_cases), len(_executable), _skipped,
    _elem_count, _er_total, _er_skipped,
)

# 无可执行用例时，pytest parametrize 空列表会产生 NOTSET 占位符
# 填入一条带 skip_reason 的哨兵用例，显示为中文跳过
if not _executable:
    _no_case = ParsedCase(
        title="无可执行用例",
        priority="",
        module="",
        steps=[],
        preconditions=[],
        expected_pages=[],
    )
    _no_case.skip_reason = "无可执行用例：所有用例均未通过筛选或元素匹配"
    _executable = [_no_case]


class TestExcelRunner:

    def setup_method(self):
        self.method = Operation_method(driver)
        self.public = Public_method()

    @pytest.mark.parametrize("case", _executable, ids=_make_case_id)
    def test_case(self, note_test_initial, case):
        """Excel 用例 — 每条独立执行，环境自动清理。"""

        # 哨兵用例：无可执行用例时的占位，直接跳过
        if case.skip_reason and not case.steps and not case.expected_pages:
            pytest.skip(case.skip_reason)

        case_id = f"R{case.row_number} [{case.priority}] {case.title}"
        logger.info(f"⏐ START  {case_id}")

        # 前置条件检查
        device_info = _get_device_info_safe()
        device_id = _get_device_id_safe()
        runtime_skip = check_conditions(case.preconditions, device_info, device_id=device_id)
        if runtime_skip:
            case.skip_reason = runtime_skip
            pytest.skip(runtime_skip)

        # 用例完整性校验（步骤号重复 / 元素缺失 / checks格式错误等）
        validation_skip = validate_case(case)
        if validation_skip:
            case.skip_reason = validation_skip
            for s in case.steps:
                s.status = "skip"
            pytest.skip(validation_skip)

        # 检查步骤（action=skip）必须配合 I 列预期结果，缺失时直接中断
        check_steps = [s for s in case.steps if s.action == "skip" and s.tag]
        if check_steps and not case.expected_pages:
            check_tags = '、'.join(s.tag for s in check_steps)
            raise AssertionError(
                f"R{case.row_number}「{case.title}」"
                f"包含 {len(check_steps)} 个检查步骤（{check_tags}），"
                f"但用例表I列（预期结果列）为空。"
                f"请在I列中添加对应的【】标记。"
            )

        # I 列为空 → WARNING（无检查步骤时的提醒）
        if not case.expected_pages:
            logger.warning(f"R{case.row_number} [{case.title}] 无预期结果，缺少断言")

        # 过滤无 element_key 的步骤（skip 除外）
        for s in case.steps:
            if not s.element_key and s.action != "skip":
                s.status = "skip"
        steps = [s for s in case.steps
                 if s.element_key or s.action == "skip"]

        # 构建 tag → 预期结果索引（支持同名 tag 多次出现时按序匹配）
        expected_by_tag: dict[str, list[ExpectedPageRef]] = {}
        for ep in case.expected_pages:
            expected_by_tag.setdefault(ep.tag, []).append(ep)
        expected_consumed: dict[str, int] = {}

        try:
            for step in steps:
                clear_element_ctx()
                set_step_context(f"R{case.row_number} 步骤{step.seq}")
                logger.info(f"  {step.raw.strip()}")
                _dispatch_step(self.method, self.public, step)
                step.status = "pass"
                clear_step_context()
                # 检查步骤 → 按 tag 关联到预期结果
                if step.action == "skip" and step.tag:
                    _check_expected_by_tag(
                        self.method, expected_by_tag, expected_consumed, step,
                        row=case.row_number, title=case.title)

            # 最终检查：I 列未消费的预期结果（DEBUG，仅调试时可见）
            unconsumed = []
            for tag, entries in expected_by_tag.items():
                used = expected_consumed.get(tag, 0)
                for i in range(used, len(entries)):
                    unconsumed.append(entries[i].raw)
            if unconsumed:
                logger.debug(
                    f"R{case.row_number} 预期结果中以下行未匹配到检查步骤:\n" +
                    "\n".join(f"  ↳ {u}" for u in unconsumed)
                )
            logger.info(f"⏐ PASSED  {case_id}")
        except Exception as e:
            for s in steps:
                if not s.status:
                    s.status = "fail"
            logger.error(f"⏐ FAILED  {case_id} — {e}")
            raise
        finally:
            clear_step_context()
            clear_element_ctx()

    @classmethod
    def teardown_class(cls):
        """全部用例执行完后，打印汇总。"""
        from boox_automation.engine.reporter import ExcelReporter
        summary = {"pass": 0, "fail": 0, "skip": 0}
        for case in _all_cases:
            status = ExcelReporter._case_status(case)
            summary[status] = summary.get(status, 0) + 1
        logger.info(f"用例汇总: 通过={summary['pass']}, 不通过={summary['fail']}, 跳过={summary['skip']}")


def _dispatch_step(method, public, step):
    """根据 action 类型分发执行。"""
    ek = step.element_key

    if step.action == "click":
        method.xpath_text_click(element_key=ek)

    elif step.action == "assert_toast":
        info = _element_matcher.get_element(ek) or {}
        toast_true = info.get("toast_true", step.tag)
        method.wait_check_toast(toast_true=toast_true, toast_timeout=5)

    elif step.action == "input":
        method.wait_input_box(element_key=ek)

    elif step.action == "long_press":
        method.wait_for_press_name(element_key=ek)

    elif step.action == "skip":
        pass  # 检查/查看/校验步骤，不需要 UI 操作

    elif step.action == "click_coord":
        method.click_by_coord(ek)

    elif step.action == "long_press_coord":
        method.long_press_by_coord(ek)

    elif step.action == "swipe_coord":
        method.swipe_by_coord(ek)

    elif step.action == "adb_cmd":
        method.execute_adb_command(ek)

    else:
        logging.warning(f"步骤{step.seq}: 【{step.tag}】未知动作类型: {step.action}")


def _check_expected_by_tag(method, expected_by_tag: dict, consumed: dict, step, *,
                           row: int = 0, title: str = "") -> None:
    """按 tag 匹配预期结果并执行检查。

    检查步骤（action=skip）通过 tag 与 I 列预期结果关联，不再依赖步骤号。
    同名 tag 按出现顺序一一对应（第一次出现的 检查【X】→ 第一个 I 列【X】）。
    """
    tag = step.tag
    entries = expected_by_tag.get(tag, [])

    if not entries:
        available_tags = sorted(expected_by_tag.keys())
        available_hint = '、'.join(available_tags) if available_tags else '(I列无可解析的预期结果)'
        case_info = f"R{row}「{title}」" if row else ""
        raise AssertionError(
            f"{case_info} 步骤{step.seq}：检查【{tag}】在用例表I列中未定义预期结果。\n"
            f"  请在I列添加【{tag}】行，当前I列已有的【】标记: {available_hint}"
        )

    idx = consumed.get(tag, 0)
    if idx >= len(entries):
        case_info = f"R{row}「{title}」" if row else ""
        raise AssertionError(
            f"{case_info} 步骤{step.seq}：检查【{tag}】的预期结果已用完"
            f"（用例表I列仅{len(entries)}条同名匹配），请检查是否多写了检查步骤"
        )

    ep = entries[idx]
    ep.step_seq = step.seq
    consumed[tag] = idx + 1

    if not ep.expected_key:
        ep.status = "fail"
        from boox_automation.engine.elements import get_element_loader
        from boox_automation.engine.diagnostics import (
            expected_sheet_not_found, expected_tag_not_matched,
        )
        loader = get_element_loader()
        diag = loader.get_expected_diagnostics()

        modules = diag.get("modules", [])
        found = diag.get("sheets_found", [])
        match_index = diag.get("match_index", {})
        total = diag.get("total_results", 0)
        skipped = diag.get("skipped", [])

        searched = [f"预期结果【{m}】" for m in modules] if modules else ["(未搜索任何模块)"]
        module_hint = (modules[1] if len(modules) > 1 else (modules[0] if modules else "模块名"))

        if total == 0 and not found:
            raise AssertionError(
                f"步骤{step.seq}：检查【{tag}】已关联到用例I列，但未加载到任何预期结果工作表。\n"
                + expected_sheet_not_found(searched, module_hint)
            )
        else:
            raise AssertionError(
                f"步骤{step.seq}：检查【{tag}】已关联到用例I列，但在预期结果工作表的 B 列中未找到匹配。\n"
                + expected_tag_not_matched(tag, found, total, match_index, module_hint, skipped)
            )

    try:
        _dispatch_expected_page(method, ep)
        ep.status = "pass"
    except Exception:
        ep.status = "fail"
        raise


def _wait_for_xml_elements(expected_xml: str, expected_key: str, step_seq: int) -> None:
    """从预期XML提取resource-id，显式等待元素就位后再获取page_source。"""
    import xml.etree.ElementTree as ET
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support import expected_conditions as EC
    from selenium.webdriver.support.wait import WebDriverWait
    from selenium.common.exceptions import TimeoutException
    from boox_automation.core.config import timeout_xml_element_wait

    try:
        root = ET.fromstring(expected_xml)
    except ET.ParseError:
        return

    from boox_automation.engine.xml_checker import _is_meaningful

    resource_ids = set()
    for elem in root.iter():
        rid = elem.get('resource-id', '')
        if rid and rid.strip() and _is_meaningful(elem):
            resource_ids.add(rid.strip())

    if not resource_ids:
        return

    timeout = timeout_xml_element_wait()
    wait = WebDriverWait(driver, timeout)
    missing_ids = []
    for rid in resource_ids:
        try:
            wait.until(EC.presence_of_element_located((By.ID, rid)))
        except TimeoutException:
            missing_ids.append(rid)

    if missing_ids:
        logger.warning(
            f"预期结果【{expected_key}】（步骤{step_seq}）"
            f"等待XML元素超时({timeout}s)，未找到{len(missing_ids)}/{len(resource_ids)}个: "
            f"{missing_ids[:3]}{'...' if len(missing_ids) > 3 else ''}"
        )
    else:
        logger.debug(
            f"预期结果【{expected_key}】（步骤{step_seq}）"
            f"XML元素就位完成 ({len(resource_ids)}个)"
        )


def _dispatch_expected_page(method, ep) -> None:
    """执行单条预期结果校验。

    断言类型由预期结果 Sheet E 列（操作）唯一决定。
    优先级: toast → D列XPath检查 → C列XML对比 → 本地文件回退
    """
    from boox_automation.engine.elements import (
        get_element_loader, _resolve_device_content,
        _load_expected_from_file, _check_elements_by_xpath,
        _check_checked_state, _EXPECTED_ACTION_MAP,
    )
    from boox_automation.engine.xml_checker import XmlChecker
    from boox_automation.driver import driver

    loader = get_element_loader()
    page_info = loader.get_expected_page(ep.expected_key)
    if not page_info:
        logger.warning(f"步骤{ep.step_seq}：预期结果【{ep.tag}】（key={ep.expected_key}）在预期结果工作表中未找到")
        ep.status = "skip"
        return

    action = page_info.get("action", "")
    check_mode = _EXPECTED_ACTION_MAP.get(action, "")
    if not check_mode:
        logger.error(f"步骤{ep.step_seq}：预期结果【{ep.tag}】（key={ep.expected_key}）E列操作无效或缺失: '{action}'")
        ep.status = "fail"
        raise AssertionError(f"预期结果 E 列操作无效: '{action}'")

    # toast 模式: 不走 XML 对比，直接轮询 page_source。toast 文本来自 B 列（匹配文本）
    if check_mode == 'toast':
        toast_text = page_info.get("match", ep.tag)
        if not method.wait_check_toast(toast_true=toast_text, toast_timeout=5):
            raise AssertionError(
                f"步骤{ep.step_seq}：预期结果【{ep.tag}】Toast'{toast_text}'未在{5}秒内出现")
        ep.status = 'pass'
        return
    elif check_mode == 'toast_not':
        toast_text = page_info.get("match", ep.tag)
        if method.wait_check_toast(toast_true=toast_text, toast_timeout=5):
            raise AssertionError(
                f"步骤{ep.step_seq}：预期结果【{ep.tag}】Toast不应出现'{toast_text}'但已检测到")
        ep.status = 'pass'
        return

    # checked / unchecked 模式：走 D 列 XPath 读取元素 checked 属性
    if check_mode in ('checked', 'unchecked'):
        element_checks = page_info.get("element_checks", "")
        if not element_checks:
            raise AssertionError(
                f"步骤{ep.step_seq}：预期结果【{ep.tag}】E列为'{action}'，"
                f"但D列（定位元素）为空。选中状态检查需提供元素定位元素。"
            )
        device_info = _get_device_info_safe()
        checks_content = _resolve_device_content(
            element_checks, device_info, key=ep.expected_key)
        if not checks_content:
            raise AssertionError(
                f"步骤{ep.step_seq}：预期结果【{ep.tag}】D列"
                f"无匹配当前设备的内容，无法执行选中状态检查。"
            )
        _check_checked_state(checks_content, mode=check_mode,
                             expected_key=ep.expected_key, step_seq=ep.step_seq)
        ep.status = "pass"
        return

    device_info = _get_device_info_safe()

    # D 列: 检查元素（元素级 XPath 检查）
    element_checks = page_info.get("element_checks", "")
    if element_checks:
        checks_content = _resolve_device_content(element_checks, device_info, key=ep.expected_key)
        if checks_content:
            try:
                _check_elements_by_xpath(checks_content, mode=check_mode,
                                         expected_key=ep.expected_key, step_seq=ep.step_seq)
                ep.status = "pass"
                return
            except AssertionError:
                ep.status = "fail"
                raise
            except Exception as e:
                logger.error(f"步骤{ep.step_seq}：预期结果【{ep.tag}】（key={ep.expected_key}）元素检查失败: {e}")
                ep.status = "fail"
                raise

    # C 列: 页面 XML
    raw_content = page_info.get("content", "")
    content = _resolve_device_content(raw_content, device_info, key=ep.expected_key)

    # C/D 均为空 → 本地文件回退
    if not content:
        file_content = _load_expected_from_file(ep.expected_key)
        if file_content:
            content = _resolve_device_content(file_content, device_info, key=ep.expected_key)

    if not content:
        logger.warning(f"步骤{ep.step_seq}：预期结果【{ep.tag}】（key={ep.expected_key}）无匹配的设备内容，跳过")
        ep.status = "skip"
        return

    # 等待预期 XML 元素加载完成
    _wait_for_xml_elements(content, ep.expected_key, ep.step_seq)

    # XML 签名对比
    try:
        actual_xml = driver.page_source
    except Exception as e:
        logger.error(f"步骤{ep.step_seq}：预期结果【{ep.tag}】（key={ep.expected_key}）获取 page_source 失败: {e}")
        ep.status = "skip"
        return

    try:
        result = XmlChecker.check(content, actual_xml, mode=check_mode)
    except Exception as e:
        logger.error(
            f"步骤{ep.step_seq}：预期结果【{ep.tag}】（key={ep.expected_key}）XML 解析失败: {e}\n"
            f"XML 前200字符: {content[:200]}"
        )
        ep.status = "skip"
        return

    ep.status = result.status
    if result.status == "fail":
        logger.info(
            f"预期结果【{ep.expected_key}】（步骤{ep.step_seq}）XML检查失败 "
            f"({result.matched_count}/{result.expected_count}):\n{result.summary()}"
        )
        raise AssertionError(result.summary())
    if result.mode == 'not_visible':
        lines = [f"预期结果【{ep.expected_key}】（步骤{ep.step_seq}）XML检查通过 (0/{result.expected_count}):"]
        for sig in result.expected:
            lines.append(f"  ↳ ✓ 不存在: {sig.to_human()}")
        logger.debug("\n".join(lines))
    else:
        lines = [f"预期结果【{ep.expected_key}】（步骤{ep.step_seq}）XML检查通过 ({result.matched_count}/{result.expected_count}):"]
        for sig in result.matched:
            lines.append(f"  ↳ ✓ 存在: {sig.to_human()}")
        for d in result.text_diffs:
            lines.append(f"  ↳  文本变更(INFO): {d['element']}")
            lines.append(f"                预期={d['expected']!r} 实际={d['actual']!r}")
        logger.debug("\n".join(lines))


if __name__ == "__main__":
    import sys
    # 确保项目根在 sys.path 中（VS Code 绿色三角形等直接 python 运行场景）
    _proj_root = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    if _proj_root not in sys.path:
        sys.path.insert(0, _proj_root)
    import pytest
    sys.exit(pytest.main([__file__, "-s"] + sys.argv[1:]))

