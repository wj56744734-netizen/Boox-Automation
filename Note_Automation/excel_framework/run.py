"""Excel 驱动测试 — PyCharm 右键运行入口。

用法:
    右键 → Run / Debug 即可启动。
    首次运行建议用干跑模式验证: 修改 DRY_RUN = True
"""
from __future__ import annotations

import logging
import sys
from pathlib import Path

from Note_Automation.Devices_list.Device_basic_information import Device_basic_information

# ============================================================
# 配置
# ============================================================
DRY_RUN = False          # True=干跑（只解析不操作设备），False=真实执行
SHEET_NAME = "笔记"       # 要执行的 Sheet
EXCEL_FILE = "test_cases.xlsx"  # excel_runner 目录下的文件名
PRIORITY_FILTER = ""     # P0=仅P0, P1=P1+P2, P2=仅P2, 空=全部
# ============================================================


from Note_Automation.excel_framework.parser import parse_steps, parse_preconditions, ParsedCase
from Note_Automation.excel_framework.matcher import ElementMatcher
from Note_Automation.excel_framework.executor import ExcelRunner
from Note_Automation.excel_framework.reporter import ExcelReporter
from Note_Automation.excel_framework.conditions import check_conditions

import openpyxl

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-5s  %(message)s",
    datefmt="%H:%M:%S",
)
logger = logging.getLogger(__name__)


def _apply_priority_filter(cases: list, spec: str) -> list:
    """按优先级筛选用例。

    spec:
        "P0" → 仅 P0
        "P1" → P1 + P2
        "P2" → 仅 P2
        ""   → 全部
    """
    if not spec:
        return cases
    spec = spec.upper().strip()
    if spec == "P0":
        return [c for c in cases if c.priority == "P0"]
    if spec == "P1":
        return [c for c in cases if c.priority in ("P1", "P2")]
    if spec == "P2":
        return [c for c in cases if c.priority == "P2"]
    return cases


def _preflight_check():
    """前置检查：设备连接 → 语言 → WiFi → 型号注册 → 测试文件。"""
    logging.getLogger('selenium').setLevel(logging.WARNING)
    logging.getLogger('urllib3').setLevel(logging.ERROR)
    logging.getLogger('appium').setLevel(logging.WARNING)

    devices = Device_basic_information()
    device_id = devices.get_connected_device_ids()

    logging.info("=" * 40)
    logging.info("  笔记自动化测试 - Excel Runner")
    logging.info("=" * 40)

    devices.check_device_language(device_id)
    devices.get_wifi(device_id)

    device_info = devices.get_device_info()
    if device_info is None:
        raise RuntimeError(
            "设备型号未注册，请检查 Devices_list/devices/ 下的 YAML 文件"
        )
    logging.info(f"设备型号: {device_info.get('device_name')} 已注册 ✓")

    devices.check_test_files(device_id)
    logging.info("-" * 40)

    return device_id, device_info


base_dir = Path(__file__).resolve().parent
excel_path = base_dir / EXCEL_FILE

if not excel_path.exists():
    print(f"文件不存在: {excel_path}")
    sys.exit(1)

# 解析
wb = openpyxl.load_workbook(str(excel_path))
if SHEET_NAME not in wb.sheetnames:
    print(f"Sheet '{SHEET_NAME}' 不存在，可选: {wb.sheetnames}")
    sys.exit(1)

ws = wb[SHEET_NAME]
matcher = ElementMatcher()
cases: list[ParsedCase] = []

for row_idx in range(2, ws.max_row + 1):
    title = ws.cell(row=row_idx, column=7).value
    steps_text = ws.cell(row=row_idx, column=10).value
    priority = ws.cell(row=row_idx, column=8).value
    module = ws.cell(row=row_idx, column=4).value
    precondition_text = ws.cell(row=row_idx, column=9).value  # I 列

    if not title or not steps_text:
        continue

    # 解析前置条件
    preconditions, variables = parse_preconditions(precondition_text)

    case = ParsedCase(
        title=str(title).strip(),
        priority=str(priority).strip() if priority else "",
        module=str(module).strip() if module else "",
        row_number=row_idx,
        steps=parse_steps(str(steps_text).strip()),
        preconditions=preconditions,
        variables=variables,
    )
    for step in case.steps:
        if step.tag:
            step.element_key = matcher.match(
                step.tag, page_context=case.module,
                case_context=f"R{case.row_number} [{case.priority}] {case.title}")
            if step.element_key:
                yaml_action = matcher.suggest_action(step.element_key)
                if yaml_action:
                    step.action = yaml_action
    cases.append(case)

# 统计
filtered_cases = _apply_priority_filter(cases, PRIORITY_FILTER)
total_steps = sum(len(c.steps) for c in cases)
tagged = sum(1 for c in cases for s in c.steps if s.tag)
matched = sum(1 for c in cases for s in c.steps if s.element_key)
executable = sum(1 for c in filtered_cases if any(s.element_key for s in c.steps))

print(f"\nSheet: {SHEET_NAME}")
print(f"用例总数: {len(cases)}  筛选后: {len(filtered_cases)}  有标记: {tagged}  已匹配: {matched}  可执行: {executable}")
if PRIORITY_FILTER:
    print(f"优先级筛选: {PRIORITY_FILTER}\n")

if DRY_RUN:
    # 干跑：展示解析结果（含前置条件）
    for case in filtered_cases:
        tags = [s for s in case.steps if s.tag]
        if not tags:
            continue
        unmatched = [s for s in tags if not s.element_key]
        flag = "⚠" if unmatched else "✓"
        print(f"{flag} [{case.priority}] {case.title}")

        # 前置条件
        if case.preconditions:
            for pc in case.preconditions:
                type_label = {"condition": "条件", "variable": "变量", "descriptive": "描述"}
                print(f"    [前置] {type_label.get(pc.type, '?')}: {pc.raw}")

        for step in case.steps:
            if step.tag:
                icon = "✓" if step.element_key else "✗"
                input_hint = f" → 【{step.input_ref}】" if step.input_ref else ""
                print(f"    {icon} 步骤{step.seq}: {step.action}【{step.tag}】{input_hint}→ {step.element_key or '未匹配'}")
        print()
else:
    # 前置检查
    try:
        device_id, device_info = _preflight_check()
    except RuntimeError as e:
        print(f"前置检查失败: {e}")
        sys.exit(1)

    # 初始化 driver
    from Note_Automation.config import ensure_driver_alive
    ensure_driver_alive()

    # 真实执行
    runner = ExcelRunner(device_info=device_info)
    reporter = ExcelReporter(str(excel_path))

    executed = 0
    skipped = 0
    for case in filtered_cases:
        if not any(s.element_key for s in case.steps):
            continue

        # 前置条件检查
        skip_reason = check_conditions(case.preconditions, device_info)
        if skip_reason:
            case.skip_reason = skip_reason
            print(f"[{case.priority}] {case.title} — 跳过: {skip_reason}")
            skipped += 1
            continue

        print(f"[{case.priority}] {case.title}")
        runner.run_case(case)
        executed += 1
        status = "✓" if all(s.status == "pass" for s in case.steps if s.element_key) else "✗"
        print(f"  {status} 步骤: {' → '.join(s.status for s in case.steps if s.element_key)}\n")

    reporter.write_results(SHEET_NAME, cases)
    reporter.save()
    print(f"\n执行完成: {executed} 个用例, 跳过 {skipped} 个，结果已写回 {EXCEL_FILE}")
