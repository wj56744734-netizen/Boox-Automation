"""结果回写器：将执行结果写回 Excel 的「测试结果」列。"""
from __future__ import annotations

import logging
from pathlib import Path

import openpyxl

from Note_Automation.excel_framework.parser import ParsedCase

logger = logging.getLogger(__name__)

# Excel 列索引（0-indexed）
COL_RESULT = 12   # M 列：测试结果
COL_REMARK = 13   # N 列：测试备注


class ExcelReporter:
    """将执行结果写回 Excel 文件。"""

    def __init__(self, filepath: str):
        self.filepath = Path(filepath)
        if not self.filepath.exists():
            raise FileNotFoundError(f"Excel 文件不存在: {filepath}")
        self._wb = openpyxl.load_workbook(filepath)

    def write_results(self, sheet_name: str, cases: list[ParsedCase]) -> None:
        """将用例执行结果写回指定 sheet。"""
        ws = self._wb[sheet_name]

        summary = {"pass": 0, "fail": 0, "skip": 0}

        for case in cases:
            row = case.row_number
            status = self._case_status(case)
            summary[status] = summary.get(status, 0) + 1

            cell = ws.cell(row=row, column=COL_RESULT + 1)  # openpyxl 1-indexed
            status_map = {
                "pass": "通过",
                "fail": "不通过",
                "skip": "跳过",
            }
            cell.value = status_map.get(status, "")

            # 测试备注：写详细步骤 + 跳过原因
            remark = self._build_detail(case)
            if remark:
                ws.cell(row=row, column=COL_REMARK + 1).value = remark

        logger.info(f"{sheet_name}: 通过={summary['pass']}, 不通过={summary['fail']}, 跳过={summary['skip']}")

    @staticmethod
    def _case_status(case: ParsedCase) -> str:
        """根据步骤状态判定用例整体结果。"""
        # 前置条件不满足 → 跳过
        if case.skip_reason:
            return "skip"

        statuses = {s.status for s in case.steps}
        # "" 表示步骤未执行（setup 失败或异常中断）→ 失败
        if "fail" in statuses or "" in statuses:
            return "fail"
        if not statuses or statuses == {"skip"}:
            return "skip"
        return "pass"

    @staticmethod
    def _build_detail(case: ParsedCase) -> str:
        """生成步骤执行详情（含前置条件跳过原因）。"""
        lines = []

        if case.skip_reason:
            lines.append(f"✗ {case.skip_reason}")

        for s in case.steps:
            if s.status == "skip" and not s.tag:
                continue
            if s.status == "skip" and case.skip_reason:
                continue  # 前置条件导致的批量 skip 不重复展示
            icon = {"pass": "✓", "fail": "✗", "skip": "-"}.get(s.status, "?")
            lines.append(f"{icon} 步骤{s.seq}: {s.action}【{s.tag}】→ {s.element_key or '未匹配'}")

        return "\n".join(lines)

    def save(self, output_path: str | None = None) -> None:
        """保存 Excel 文件。"""
        target = output_path or self.filepath
        self._wb.save(str(target))
        logger.info(f"结果已保存: {target}")
