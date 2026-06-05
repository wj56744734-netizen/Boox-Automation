"""Excel 驱动测试 — 集成到 pytest 框架。

每条 Excel 用例作为独立的 pytest 用例执行，自动获得 note_test_initial 的完整环境清理。

运行方式:
    pytest Note_Automation/Test_local_notes/test_excel_runner.py -s
"""
from __future__ import annotations

import logging
import re
from pathlib import Path

import allure
import pytest
import openpyxl

from Note_Automation.conftest import note_mark_china, note_mark_increment
from Note_Automation.config import driver
from Note_Automation.Note_class.Note_class import Operation_method
from Note_Automation.Test_local_notes.Public_method import Public_method
from Note_Automation.excel_framework.parser import parse_steps, parse_preconditions, ParsedCase
from Note_Automation.excel_framework.matcher import ElementMatcher
from Note_Automation.excel_framework.conditions import check_conditions
from Note_Automation.Note_class.Note_element.element_loader import LOCATOR_TYPE_MAP

logger = logging.getLogger(__name__)


def _resolve_checks(checks: list) -> list[dict]:
    """将 matcher 返回的原始 checks 转为 check_multi_elements 可用格式。"""
    result = []
    for c in checks:
        loc = c["locator"]
        loc_type, loc_value = loc[0], loc[1]
        by = LOCATOR_TYPE_MAP.get(loc_type.lower(), LOCATOR_TYPE_MAP["xpath"])
        result.append({"locator": (by, loc_value), "text": c.get("text", "")})
    return result

# ============================================================
SHEET_NAME = "笔记"
EXCEL_FILE = str(Path(__file__).resolve().parent.parent / "excel_framework" / "test_cases.xlsx")
PRIORITY_FILTER = "P0"   # P0=仅P0, P1=P1+P2, P2=仅P2, 空=全部
_PRIORITY_ORDER = {"P0": 0, "P1": 1, "P2": 2}
# ============================================================


def _get_device_info_safe() -> dict:
    """安全获取设备信息，失败返回空字典（不影响模块加载）。"""
    try:
        from Note_Automation.Devices_list.Device_basic_information import Device_basic_information
        dbi = Device_basic_information()
        return dbi.get_device_info() or {}
    except Exception:
        return {}


_element_matcher = ElementMatcher()


def _load_cases(excel_path: str, sheet: str, priority_spec: str) -> list[ParsedCase]:
    """加载 Excel 并解析匹配，返回可执行用例列表。"""
    wb = openpyxl.load_workbook(excel_path)
    ws = wb[sheet]
    cases = []

    for row_idx in range(2, ws.max_row + 1):
        title = ws.cell(row=row_idx, column=7).value
        steps_text = ws.cell(row=row_idx, column=10).value
        priority = ws.cell(row=row_idx, column=8).value
        module = ws.cell(row=row_idx, column=4).value
        precondition_text = ws.cell(row=row_idx, column=9).value  # I 列

        if not title or not steps_text:
            continue

        # 解析前置条件
        preconditions, variables = parse_preconditions(precondition_text)

        case = ParsedCase(
            title=str(title).strip(),
            priority=str(priority).strip() if priority else "",
            module=str(module).strip() if module else "",
            row_number=row_idx,
            steps=parse_steps(str(steps_text).strip()),
            preconditions=preconditions,
            variables=variables,
        )

        ctx = f"R{row_idx} [{priority}] {title}"
        for step in case.steps:
            if step.tag:
                step.element_key = _element_matcher.match(step.tag, page_context=case.module, case_context=ctx, warn=False)
                if step.element_key:
                    step.action = _element_matcher.suggest_action(step.element_key)

        cases.append(case)

    # 在过滤前报告所有未匹配标记（全量扫描，不漏被过滤优先级的用例）
    _unmatched = []
    for case in cases:
        for step in case.steps:
            if step.tag and not step.element_key:
                _unmatched.append(f"  R{case.row_number} [{case.priority}] {case.title}: 步骤{step.seq}【{step.tag}】未匹配")
    if _unmatched:
        logger.warning(f"{len(_unmatched)} 个标记未匹配到元素:")
        for msg in _unmatched:
            logger.warning(msg)

    # 优先级筛选
    spec = priority_spec.upper().strip()
    if spec == "P0":
        cases = [c for c in cases if c.priority == "P0"]
    elif spec == "P1":
        cases = [c for c in cases if c.priority in ("P1", "P2")]
    elif spec == "P2":
        cases = [c for c in cases if c.priority == "P2"]

    return cases


def _make_case_id(case: ParsedCase) -> str:
    prefix = "[跳过]" if case.skip_reason else ""
    return f"{prefix}R{case.row_number}-{case.title[:30]}"


# 模块加载时解析 Excel，生成 parametrize 参数
_all_cases = _load_cases(EXCEL_FILE, SHEET_NAME, PRIORITY_FILTER)
_executable = [c for c in _all_cases if any(s.element_key for s in c.steps)]
_executable.sort(key=lambda c: (_PRIORITY_ORDER.get(c.priority, 99), c.row_number))


@allure.feature("Excel驱动测试")
class TestExcelRunner:

    def setup_method(self):
        self.method = Operation_method(driver)
        self.public = Public_method()

    @pytest.mark.parametrize("case", _executable, ids=_make_case_id)
    def test_case(self, note_test_initial, case):
        """Excel 用例 — 每条独立执行，环境自动清理。"""
        # 运行时前置条件检查（设备已确保连接，获取真实设备信息）
        device_info = _get_device_info_safe()
        runtime_skip = check_conditions(case.preconditions, device_info)
        if runtime_skip:
            case.skip_reason = runtime_skip
            logger.info(f"SKIPPED  R{case.row_number} [{case.priority}] {case.title} — {runtime_skip}")
            pytest.skip(runtime_skip)

        # 无 element_key 的步骤标记为跳过（非失败）
        for s in case.steps:
            if s.tag and not s.element_key:
                s.status = "skip"
        steps = [s for s in case.steps if s.element_key]

        case_id = f"R{case.row_number} [{case.priority}] {case.title}"
        try:
            with allure.step(case_id):
                for step in steps:
                    _dispatch_step(self.method, self.public, step, case.variables)
                    step.status = "pass"
            logger.info(f"PASSED   {case_id}")
        except Exception:
            for s in steps:
                if not s.status:
                    s.status = "fail"
            logger.error(f"FAILED   {case_id}")
            raise

    @classmethod
    def teardown_class(cls):
        """全部用例执行完后，打印汇总。"""
        from Note_Automation.excel_framework.reporter import ExcelReporter
        summary = {"pass": 0, "fail": 0, "skip": 0}
        for case in _all_cases:
            status = ExcelReporter._case_status(case)
            summary[status] = summary.get(status, 0) + 1
        logger.info(f"{SHEET_NAME}: 通过={summary['pass']}, 不通过={summary['fail']}, 跳过={summary['skip']}")


def _dispatch_step(method, public, step, variables: dict[str, str] | None = None):
    """根据 action 类型分发执行。"""
    ek = step.element_key
    vars_ = variables or {}

    if step.action == "click":
        method.xpath_text_click(element_key=ek)

    elif step.action in ("assert", "assert_visible"):
        info = _element_matcher.get_element(ek) or {}
        checks = info.get("checks")
        if checks:
            method.check_multi_elements(_resolve_checks(checks), element_key=ek)
        else:
            method.xpath_text_click(element_key=ek, should_click=False)

    elif step.action == "assert_not":
        try:
            method.xpath_text_click(element_key=ek, should_click=False)
            raise AssertionError(f"元素【{step.tag}】应不存在但实际存在")
        except AssertionError:
            raise
        except Exception:
            pass

    elif step.action == "assert_text":
        info = _element_matcher.get_element(ek) or {}
        expected = info.get("assert_text", step.tag)
        actual = method.obtain_element_text(element_key=ek)
        if actual != expected:
            raise AssertionError(
                f"文本不匹配: 期望='{expected}', 实际='{actual}'"
            )

    elif step.action == "assert_toast":
        info = _element_matcher.get_element(ek) or {}
        toast_true = info.get("toast_true", step.tag)
        method.wait_check_toast(toast_true=toast_true, toast_timeout=5)

    elif step.action == "input":
        # 确定输入内容：优先变量引用 → 步骤文本解析
        input_text = ""
        if step.input_ref:
            input_text = vars_.get(step.input_ref, step.input_ref)
        if not input_text:
            text_match = re.search(r"输入(.+?)(?:字符|$)", step.raw)
            input_text = text_match.group(1) if text_match else "test"
        method.wait_input_box(element_key=ek, name=input_text)

    elif step.action == "long_press":
        method.wait_for_press_name(element_key=ek)

    elif step.action == "dismiss":
        info = _element_matcher.get_element(ek) or {}
        dismiss_key = info.get("dismiss_with", "")
        checks = info.get("checks")
        try:
            if checks:
                method.check_multi_elements(_resolve_checks(checks), element_key=ek)
            else:
                method.xpath_text_click(element_key=ek, should_click=False)
            if dismiss_key:
                method.xpath_text_click(element_key=dismiss_key)
            else:
                method.xpath_text_click(element_key=ek)
        except Exception:
            pass

    else:
        logging.warning(f"未知动作类型: {step.action}")
